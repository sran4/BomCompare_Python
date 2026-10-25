#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
BOM vs SMT Feeder Setup Sheet Comparator
========================================

Compares the reference designators in a Bill of Materials (Excel) against the
designators found in an SMT feeder / setup sheet (TXT, LOG or PDF) and reports:

  * Missing in SMT   - designator is in the BOM but not on the setup sheet
  * Missing in BOM   - designator is on the setup sheet but not in the BOM
  * MPN Mismatch     - designator is in both, but the BOM MPN cannot be found on
                       the setup-sheet line(s) where the designator appears
  * Duplicate in BOM - the same designator is listed on more than one BOM line
  * Duplicate in SMT - the same designator is loaded with different parts

Requirements
------------
    pip install PySide6 pandas openpyxl xlrd pdfplumber pypdf

    * xlrd is only needed for legacy .xls files
    * pdfplumber (preferred) or pypdf is needed for PDF setup sheets

Run
---
    python bom_smt_comparator.py
"""

from __future__ import annotations

import datetime as _dt
import os
import re
import sys
import traceback
from collections import defaultdict
from dataclasses import dataclass, field
from typing import Callable, Optional

import pandas as pd

from PySide6.QtCore import (
    QAbstractTableModel,
    QModelIndex,
    QObject,
    QSettings,
    QSortFilterProxyModel,
    Qt,
    QThread,
    Signal,
    Slot,
)
from PySide6.QtGui import QColor, QFont, QKeySequence, QPalette, QShortcut
from PySide6.QtWidgets import (
    QAbstractItemView,
    QApplication,
    QCheckBox,
    QComboBox,
    QFileDialog,
    QFormLayout,
    QFrame,
    QGroupBox,
    QHBoxLayout,
    QHeaderView,
    QLabel,
    QLineEdit,
    QMainWindow,
    QMessageBox,
    QProgressBar,
    QPushButton,
    QScrollArea,
    QSpinBox,
    QSplitter,
    QTableView,
    QVBoxLayout,
    QWidget,
)

# Optional PDF back-ends -------------------------------------------------------
try:
    import pdfplumber  # type: ignore
except ImportError:  # pragma: no cover
    pdfplumber = None

try:
    from pypdf import PdfReader  # type: ignore
except ImportError:  # pragma: no cover
    PdfReader = None


# =============================================================================
# Constants
# =============================================================================
APP_NAME = "BOM ⇄ SMT Setup Comparator"
APP_ORG = "SMTTools"
APP_VERSION = "1.1.0"

SMT_EXTENSIONS = (".txt", ".log", ".pdf")
BOM_EXTENSIONS = (".xlsx", ".xls", ".xlsm")

DEFAULT_PREFIXES = "R, C, L, U, D, Q, J, P, Y, X, F, FB, FL, K, T, TP, SW, LED, BT, CN, CR, VR, RN, RP, IC, JP, ZD"
NONE_OPTION = "(none)"

# Package size codes that look like designators (e.g. "C0603", "R0402")
PACKAGE_CODES = {
    "0100", "0201", "0402", "0603", "0805", "1008", "1206", "1210",
    "1806", "1812", "2010", "2220", "2512",
}
MAX_RANGE_EXPANSION = 1000

GENERIC_REF_RE = re.compile(r"^([A-Z]{1,5})(\d{1,6})([A-Z]?)$")
BOM_RANGE_RE = re.compile(r"^([A-Z]{1,5}\d{1,6})[-~]([A-Z]{0,5}\d{1,6})$")
BOM_SPLIT_RE = re.compile(r"[,;\s/|]+")
SMT_TOKEN_SPLIT_RE = re.compile(r"[\s,;|]+")
VALUE_RE = re.compile(r"^[\d.]+[A-Z]{0,4}%?$|^\d+[RKMUNP]\d*[A-Z]{0,3}$")
DNP_RE = re.compile(
    r"\b(DNP|DNI|DNF|DNL|NOPOP|NO\s*POP|NOT\s+FITTED|NOT\s+POPULATED|NOT\s+MOUNTED|NO\s+LOAD|"
    r"DO\s+NOT\s+(?:POPULATE|PLACE|INSTALL|FIT|MOUNT|STUFF))\b"
)

# Discrepancy types
MISSING_SMT = "Missing in SMT"
MISSING_BOM = "Missing in BOM"
MPN_MISMATCH = "MPN Mismatch"
DUP_BOM = "Duplicate in BOM"
DUP_SMT = "Duplicate in SMT"
ALL_TYPES = [MISSING_SMT, MISSING_BOM, MPN_MISMATCH, DUP_BOM, DUP_SMT]

TYPE_COLORS = {
    MISSING_SMT: "#FFC7CE",   # light red
    MISSING_BOM: "#FFDADA",   # lighter red
    MPN_MISMATCH: "#FFF2CC",  # light yellow
    DUP_BOM: "#FCE4D6",       # light orange
    DUP_SMT: "#FCE4D6",
}
DARK_TYPE_COLORS = {
    MISSING_SMT: "#6E3038",
    MISSING_BOM: "#5C3338",
    MPN_MISMATCH: "#6B5520",
    DUP_BOM: "#6B4530",
    DUP_SMT: "#6B4530",
}

COL_REF = "Reference Designator"
COL_TYPE = "Discrepancy Type"
COL_BOM_MPN = "BOM MPN"
COL_SMT_MPN = "SMT MPN"
COL_NOTES = "Notes"
COL_BOM_DESC = "BOM Description"
COL_SMT_SRC = "SMT Source"

TABLE_COLUMNS = [COL_REF, COL_TYPE, COL_BOM_MPN, COL_SMT_MPN, COL_NOTES]
EXPORT_COLUMNS = TABLE_COLUMNS + [COL_BOM_DESC, COL_SMT_SRC]

SORT_ROLE = Qt.ItemDataRole.UserRole + 1


# =============================================================================
# Exceptions
# =============================================================================
class AppError(Exception):
    """Base class for errors that are shown to the user as a friendly pop-up."""

    title = "Error"

    def __init__(self, message: str, title: Optional[str] = None):
        super().__init__(message)
        if title:
            self.title = title


class FileReadError(AppError):
    title = "File Read Error"


class ColumnError(AppError):
    title = "Column Not Found"


class ParseError(AppError):
    title = "Parsing Error"


# =============================================================================
# Data classes
# =============================================================================
@dataclass
class BomEntry:
    ref: str
    mpn: str
    desc: str
    row: int
    inherited: bool = False


@dataclass
class BomData:
    entries: dict = field(default_factory=dict)          # ref -> list[BomEntry]
    dnp_refs: dict = field(default_factory=dict)         # ref -> excel row
    invalid_tokens: list = field(default_factory=list)   # (excel row, token)
    header_row: int = 1
    columns: dict = field(default_factory=dict)
    rows_with_refs: int = 0


@dataclass
class SmtEntry:
    ref: str
    mpn: str
    location: str
    context: str
    inherited: bool = False


@dataclass
class CompareOptions:
    smt_path: str
    bom_path: str
    bom_sheet: object
    header_row: int                 # 1-based; 0 = auto-detect
    ref_col: str
    mpn_col: str
    desc_col: str
    prefixes: list
    ignore_prefixes: set
    learn_prefixes: bool = True
    check_mpn: bool = True
    exclude_dnp: bool = True
    smt_column_mode: bool = False
    smt_ref_col: int = 6
    smt_mpn_col: int = 3


@dataclass
class CompareResult:
    rows: list
    stats: dict
    warnings: list
    header_row: int


# =============================================================================
# Generic helpers
# =============================================================================
def cell_text(value) -> str:
    """Convert any Excel cell value to a clean string."""
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).replace("\r", " ").replace("\n", " ").strip()


def normalize_mpn(text: str) -> str:
    """Uppercase and strip everything except A-Z / 0-9 for robust MPN matching."""
    return re.sub(r"[^A-Z0-9]", "", str(text or "").upper())


def ref_prefix(ref: str) -> str:
    m = GENERIC_REF_RE.match(ref)
    return m.group(1) if m else ref


def natural_key(ref: str):
    m = GENERIC_REF_RE.match(ref)
    if m:
        return (m.group(1), int(m.group(2)), m.group(3))
    return (ref, 0, "")


def natural_sort_string(ref: str) -> str:
    p, n, s = natural_key(ref)
    return f"{p}{n:09d}{s}"


def unique(items) -> list:
    seen, out = set(), []
    for it in items:
        if it and it not in seen:
            seen.add(it)
            out.append(it)
    return out


def col_letter(index: int) -> str:
    """0-based column index -> Excel letters."""
    letters = ""
    index += 1
    while index:
        index, rem = divmod(index - 1, 26)
        letters = chr(65 + rem) + letters
    return letters


def letters_to_index(letters: str) -> int:
    idx = 0
    for ch in letters.upper():
        idx = idx * 26 + (ord(ch) - 64)
    return idx - 1


def parse_prefix_list(text: str) -> list:
    return unique(p.strip().upper() for p in re.split(r"[,;\s]+", text or "") if p.strip().isalpha())


def expand_range(start: str, end: str) -> Optional[list]:
    """Expand 'R1'-'R5' or 'R1'-'5' into R1..R5. Returns None if not a valid range."""
    ms = GENERIC_REF_RE.match(start)
    if not ms:
        return None
    pre, num_s, suf = ms.group(1), ms.group(2), ms.group(3)
    if end.isdigit():
        e_pre, e_num_s, e_suf = pre, end, ""
    else:
        me = GENERIC_REF_RE.match(end)
        if not me:
            return None
        e_pre, e_num_s, e_suf = me.group(1), me.group(2), me.group(3)
    num, e_num = int(num_s), int(e_num_s)
    if e_pre != pre or suf or e_suf or e_num < num or (e_num - num) > MAX_RANGE_EXPANSION:
        return None
    width = len(num_s) if num_s.startswith("0") and len(num_s) > 1 else 0
    return [f"{pre}{str(i).zfill(width) if width else i}" for i in range(num, e_num + 1)]


def split_bom_designators(text: str):
    """
    Clean and split a BOM designator cell such as "R7, R45,C2" or "R1-R4; C9".
    Returns (designators, unparseable_tokens).
    """
    t = str(text or "").upper()
    t = t.replace("\u2013", "-").replace("\u2014", "-").replace("\u00a0", " ")
    t = re.sub(r"\s*([-~])\s*", r"\1", t)          # "R1 - R4" -> "R1-R4"
    t = re.sub(r"[\"'()\[\]{}]", " ", t)
    refs, bad = [], []
    for tok in BOM_SPLIT_RE.split(t):
        tok = tok.strip(".:-~")
        if not tok:
            continue
        if GENERIC_REF_RE.match(tok):
            refs.append(tok)
            continue
        m = BOM_RANGE_RE.match(tok)
        if m:
            expanded = expand_range(m.group(1), m.group(2))
            if expanded:
                refs.extend(expanded)
                continue
        bad.append(tok)
    return refs, bad


# =============================================================================
# BOM (Excel) parsing
# =============================================================================
def _is_ref_header(h: str) -> bool:
    h = h.strip().lower()
    return any(k in h for k in ("designator", "refdes", "ref des", "ref. des", "reference")) or h in {
        "ref", "refs", "ref.", "location", "locations", "position", "positions", "ref no", "ref no."
    }


def _is_mpn_header(h: str, strong: bool) -> bool:
    h = h.strip().lower()
    strong_keys = ("mpn", "manufacturer part", "mfr part", "mfg part", "mfr p/n", "mfg p/n",
                   "manufacturer p/n", "mfr. part", "mfr pn", "mfg pn", "manufacturer pn")
    if strong:
        return any(k in h for k in strong_keys)
    return h in {"part number", "part no", "part no.", "p/n", "pn", "part #", "part"} or "part number" in h


def _is_desc_header(h: str) -> bool:
    h = h.strip().lower()
    return "description" in h or h in {"desc", "desc.", "comment", "value"}


def guess_column(headers: list, kind: str) -> Optional[str]:
    if kind == "ref":
        return next((h for h in headers if _is_ref_header(h)), None)
    if kind == "mpn":
        return (next((h for h in headers if _is_mpn_header(h, True)), None)
                or next((h for h in headers if _is_mpn_header(h, False)), None))
    if kind == "desc":
        return next((h for h in headers if _is_desc_header(h)), None)
    return None


def list_bom_sheets(path: str) -> list:
    try:
        with pd.ExcelFile(path) as xf:
            return [str(s) for s in xf.sheet_names]
    except FileNotFoundError:
        raise FileReadError(f"BOM file not found:\n{path}")
    except PermissionError:
        raise FileReadError(f"Permission denied while opening:\n{path}\n\nIs the file open in another program?")
    except ImportError as e:
        raise FileReadError(
            f"A required Excel engine is not installed:\n{e}\n\n"
            "Install it with:  pip install openpyxl xlrd"
        )
    except Exception as e:
        raise FileReadError(f"Unable to open the BOM workbook. It may be corrupt or not an Excel file.\n\nDetails: {e}")


def read_bom_raw(path: str, sheet, nrows: Optional[int] = None) -> pd.DataFrame:
    try:
        return pd.read_excel(path, sheet_name=sheet if sheet not in (None, "") else 0,
                             header=None, dtype=object, nrows=nrows)
    except FileNotFoundError:
        raise FileReadError(f"BOM file not found:\n{path}")
    except PermissionError:
        raise FileReadError(f"Permission denied while reading:\n{path}\n\nIs the file open in another program?")
    except ImportError as e:
        raise FileReadError(
            f"A required Excel engine is not installed:\n{e}\n\n"
            "Install it with:  pip install openpyxl xlrd"
        )
    except ValueError as e:
        raise FileReadError(f"Could not read sheet '{sheet}' from the BOM.\n\nDetails: {e}")
    except Exception as e:
        raise FileReadError(f"Unable to read the BOM file. It may be corrupt or password protected.\n\nDetails: {e}")


def detect_header_row(raw: pd.DataFrame, scan_rows: int = 40) -> int:
    """Return the 0-based index of the most likely header row."""
    best_idx, best_score = 0, 0
    for i in range(min(scan_rows, len(raw))):
        cells = [cell_text(v) for v in raw.iloc[i].tolist()]
        cells = [c for c in cells if c]
        if not cells:
            continue
        has_ref = any(_is_ref_header(c) for c in cells)
        score = (3 if has_ref else 0)
        score += 2 if any(_is_mpn_header(c, True) or _is_mpn_header(c, False) for c in cells) else 0
        score += 1 if any(_is_desc_header(c) for c in cells) else 0
        if has_ref and score >= 4:
            return i
        if score > best_score:
            best_idx, best_score = i, score
    return best_idx


def build_headers(raw: pd.DataFrame, header_idx: int) -> list:
    names, seen = [], {}
    for j, v in enumerate(raw.iloc[header_idx].tolist()):
        name = cell_text(v) or f"Column {col_letter(j)}"
        if name in seen:
            seen[name] += 1
            name = f"{name} ({seen[name]})"
        else:
            seen[name] = 1
        names.append(name)
    return names


def peek_bom_headers(path: str, sheet, header_row: int):
    """Read only the top of the sheet to populate the column selectors."""
    nrows = max(60, header_row + 1)
    raw = read_bom_raw(path, sheet, nrows=nrows)
    if raw.empty or raw.dropna(how="all").empty:
        raise ParseError(f"The sheet '{sheet}' appears to be empty.")
    idx = header_row - 1 if header_row > 0 else detect_header_row(raw)
    if idx >= len(raw):
        raise ColumnError(f"Header row {idx + 1} is beyond the end of the sheet.")
    return build_headers(raw, idx), idx


def resolve_column(headers: list, spec: str, label: str, required: bool) -> Optional[int]:
    """
    Resolve a column given as a header name (case-insensitive), a 1-based index ("3"),
    or an Excel column letter ("C").
    """
    spec = (spec or "").strip()
    if not spec or spec == NONE_OPTION:
        if required:
            raise ColumnError(f"No column was specified for '{label}'.")
        return None

    lower_map = {}
    for i, h in enumerate(headers):
        lower_map.setdefault(h.strip().lower(), i)
    if spec.lower() in lower_map:
        return lower_map[spec.lower()]

    if spec.isdigit():
        i = int(spec) - 1
        if 0 <= i < len(headers):
            return i

    if re.fullmatch(r"[A-Za-z]{1,3}", spec):
        i = letters_to_index(spec)
        if 0 <= i < len(headers):
            return i

    partial = [i for i, h in enumerate(headers) if spec.lower() in h.lower()]
    if len(partial) == 1:
        return partial[0]

    available = "\n".join(f"  {col_letter(i)} ({i + 1}): {h}" for i, h in enumerate(headers))
    raise ColumnError(
        f"Could not find the '{label}' column using \"{spec}\".\n\n"
        f"Enter a header name, a 1-based column number, or an Excel column letter.\n\n"
        f"Available columns:\n{available}"
    )


def load_bom(opts: CompareOptions) -> BomData:
    raw = read_bom_raw(opts.bom_path, opts.bom_sheet)
    if raw.empty or raw.dropna(how="all").empty:
        raise ParseError("The selected BOM sheet is empty.")

    header_idx = opts.header_row - 1 if opts.header_row > 0 else detect_header_row(raw)
    if header_idx >= len(raw):
        raise ColumnError(f"Header row {header_idx + 1} is beyond the end of the sheet ({len(raw)} rows).")

    headers = build_headers(raw, header_idx)
    ref_i = resolve_column(headers, opts.ref_col, "Reference Designators", True)
    mpn_i = resolve_column(headers, opts.mpn_col, "Manufacturer Part Number", True)
    desc_i = resolve_column(headers, opts.desc_col, "Description", False)
    if ref_i == mpn_i:
        raise ColumnError("The Reference Designator and MPN selections point to the same column.")

    data = BomData(
        header_row=header_idx + 1,
        columns={"ref": headers[ref_i], "mpn": headers[mpn_i],
                 "desc": headers[desc_i] if desc_i is not None else ""},
    )

    prev_mpn = prev_desc = ""
    for offset, values in enumerate(raw.iloc[header_idx + 1:].itertuples(index=False, name=None)):
        excel_row = header_idx + 2 + offset
        cells = [cell_text(v) for v in values]
        ref_txt = cells[ref_i] if ref_i < len(cells) else ""
        mpn = cells[mpn_i] if mpn_i < len(cells) else ""
        desc = cells[desc_i] if desc_i is not None and desc_i < len(cells) else ""

        if not ref_txt:
            if mpn or desc:
                prev_mpn, prev_desc = mpn, desc
            continue

        # Continuation row: designators spilled onto a row with nothing else on it
        inherited = False
        if not mpn and not desc and prev_mpn:
            others = [c for j, c in enumerate(cells) if j != ref_i and c]
            if not others:
                mpn, desc, inherited = prev_mpn, prev_desc, True

        refs, bad = split_bom_designators(ref_txt)
        data.invalid_tokens.extend((excel_row, b) for b in bad if not DNP_RE.search(b))
        if not refs:
            continue
        data.rows_with_refs += 1

        if opts.exclude_dnp and any(DNP_RE.search(c.upper()) for c in cells if c):
            for r in refs:
                data.dnp_refs.setdefault(r, excel_row)
            continue

        for r in refs:
            data.entries.setdefault(r, []).append(BomEntry(r, mpn, desc, excel_row, inherited))
        if mpn:
            prev_mpn, prev_desc = mpn, desc

    if not data.entries and not data.dnp_refs:
        raise ParseError(
            f"No reference designators were found in column '{headers[ref_i]}' "
            f"(header row {header_idx + 1}).\n\n"
            "Check that the correct sheet, header row and designator column are selected."
        )
    return data


# =============================================================================
# SMT setup sheet parsing
# =============================================================================
class SmtScanner:
    """Finds reference designators and guesses part numbers in free-form setup sheet text."""

    def __init__(self, prefixes):
        clean = sorted({p.strip().upper() for p in prefixes if p and p.strip().isalpha()},
                       key=len, reverse=True)
        if not clean:
            raise AppError("No valid designator prefixes were configured for the SMT scan.",
                           "Configuration Error")
        self.prefixes = set(clean)
        alt = "|".join(map(re.escape, clean))
        self._scan = re.compile(
            rf"(?<![A-Z0-9_.\-])((?:{alt})\d{{1,6}}[A-Z]?)"
            rf"(?:[-~]((?:{alt})?\d{{1,6}}))?"
            rf"(?![A-Z0-9_\-])"
        )

    @staticmethod
    def _valid(ref: str) -> bool:
        m = GENERIC_REF_RE.match(ref)
        return bool(m) and m.group(2) not in PACKAGE_CODES

    def find(self, line_upper: str) -> list:
        refs = []
        for m in self._scan.finditer(line_upper):
            start, end = m.group(1), m.group(2)
            if end:
                expanded = expand_range(start, end) if self._valid(start) else None
                if expanded:
                    refs.extend(expanded)
                    continue
                candidates = [start] + ([] if end.isdigit() else [end])
            else:
                candidates = [start]
            refs.extend(c for c in candidates if self._valid(c))
        return refs

    def is_designator_token(self, token: str) -> bool:
        return bool(self._scan.fullmatch(token))

    def guess_mpn(self, tokens: list, bom_mpn_norms: set) -> Optional[str]:
        # 1) Exact hit against a known BOM MPN is the strongest signal
        for tok in tokens:
            n = normalize_mpn(tok)
            if len(n) >= 4 and n in bom_mpn_norms:
                return tok
        # 2) Heuristic: longest alphanumeric token that is not a designator / value
        best, best_len = None, 0
        for tok in tokens:
            if self.is_designator_token(tok) or VALUE_RE.match(tok):
                continue
            n = normalize_mpn(tok)
            if len(n) < 6 or not re.search(r"[A-Z]", n) or not re.search(r"\d", n):
                continue
            if len(n) > best_len:
                best, best_len = tok, len(n)
        return best


def _decode_text(raw: bytes) -> str:
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16")
    if b"\x00" in raw[:2000]:
        for enc in ("utf-16-le", "utf-16-be"):
            try:
                return raw.decode(enc)
            except UnicodeDecodeError:
                pass
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            pass
    return raw.decode("latin-1", errors="replace")


def _read_pdf_pages(path: str) -> list:
    if pdfplumber is None and PdfReader is None:
        raise FileReadError("PDF support requires 'pdfplumber' or 'pypdf'.\n\nInstall with:  pip install pdfplumber pypdf")

    errors = []
    if pdfplumber is not None:
        try:
            with pdfplumber.open(path) as pdf:
                pages = [(i, page.extract_text() or "") for i, page in enumerate(pdf.pages, 1)]
            if any(t.strip() for _, t in pages):
                return pages
        except Exception as e:
            errors.append(f"pdfplumber: {e}")

    if PdfReader is not None:
        try:
            reader = PdfReader(path)
            if reader.is_encrypted:
                try:
                    ok = reader.decrypt("")
                except Exception:
                    ok = 0
                if not ok:
                    raise FileReadError("The PDF is password protected. Remove the password and try again.")
            pages = [(i, page.extract_text() or "") for i, page in enumerate(reader.pages, 1)]
            if any(t.strip() for _, t in pages):
                return pages
        except FileReadError:
            raise
        except Exception as e:
            errors.append(f"pypdf: {e}")

    if errors:
        raise FileReadError("The PDF could not be read. It may be corrupt or protected.\n\nDetails:\n" + "\n".join(errors))
    raise ParseError(
        "The PDF contains no extractable text. It is probably a scanned image.\n\n"
        "Run it through OCR (e.g. 'Save as searchable PDF') or export the setup sheet as text from the SMT software."
    )


def read_smt_lines(path: str) -> list:
    """Return a list of (location, line_text) tuples."""
    if not os.path.isfile(path):
        raise FileReadError(f"SMT setup sheet not found:\n{path}")
    ext = os.path.splitext(path)[1].lower()
    lines = []
    try:
        if ext == ".pdf":
            for page_no, text in _read_pdf_pages(path):
                for ln, line in enumerate(text.splitlines(), 1):
                    lines.append((f"Page {page_no}, line {ln}", line))
        elif ext in (".txt", ".log"):
            with open(path, "rb") as fh:
                text = _decode_text(fh.read())
            lines = [(f"Line {i}", l) for i, l in enumerate(text.splitlines(), 1)]
        else:
            raise FileReadError(f"Unsupported SMT file type '{ext}'. Use .txt, .log or .pdf.")
    except PermissionError:
        raise FileReadError(f"Permission denied while reading:\n{path}\n\nIs the file open in another program?")
    except OSError as e:
        raise FileReadError(f"Unable to read the SMT setup sheet.\n\nDetails: {e}")

    if not any(l.strip() for _, l in lines):
        raise ParseError("The SMT setup sheet is empty.")
    return lines


def parse_smt(lines: list, scanner: SmtScanner, bom_mpn_norms: set) -> dict:
    """
    Build ref -> list[SmtEntry]. Lines that contain designators but no part number
    inherit the part number from the nearest preceding line that had one
    (handles wrapped designator lists and 'part line + designator line' layouts).
    """
    data = defaultdict(list)
    last_mpn, last_ctx, last_idx = None, "", -999
    for idx, (loc, raw_line) in enumerate(lines):
        line = raw_line.upper().strip()
        if not line:
            continue
        refs = scanner.find(line)
        tokens = [t.strip("\"'()[]{}:") for t in SMT_TOKEN_SPLIT_RE.split(line)]
        tokens = [t for t in tokens if t]
        mpn = scanner.guess_mpn(tokens, bom_mpn_norms)
        ctx = raw_line.strip()
        inherited = False

        if mpn:
            last_mpn, last_ctx, last_idx = mpn, ctx, idx
        elif refs and last_mpn and idx - last_idx <= 3:
            mpn, inherited = last_mpn, True
            ctx = f"{last_ctx}  ⏎  {ctx}"
            last_idx = idx  # keep long wrapped lists attached to the same part

        for r in refs:
            data[r].append(SmtEntry(r, mpn or "", loc, ctx, inherited))
    return dict(data)


def parse_smt_columns(lines: list, scanner: SmtScanner, bom_mpn_norms: set,
                      ref_col: int, mpn_col: int) -> dict:
    """Parse a machine dump split on tabs, 2+ spaces, or pipes, using 1-based columns."""
    if ref_col < 1:
        raise AppError("Setup RefDes column must be 1 or greater.", "Configuration Error")
    data = defaultdict(list)
    last_mpn, last_ctx, last_idx = None, "", -999
    for idx, (loc, raw_line) in enumerate(lines):
        raw = raw_line.strip()
        if not raw:
            continue
        parts = re.split(r"\t+|\s{2,}|\|", raw)
        if ref_col > len(parts):
            continue
        refs = scanner.find(parts[ref_col - 1].upper())
        if not refs:
            continue
        mpn = ""
        if 1 <= mpn_col <= len(parts):
            cand = parts[mpn_col - 1].strip()
            if cand and not scanner.find(cand.upper()):
                mpn = cand
        if not mpn:
            tokens = [t.strip("\"'()[]{}:") for t in SMT_TOKEN_SPLIT_RE.split(raw.upper()) if t]
            tokens = [t for t in tokens if t]
            mpn = scanner.guess_mpn(tokens, bom_mpn_norms) or ""
        inherited = False
        ctx = raw
        if mpn:
            last_mpn, last_ctx, last_idx = mpn, ctx, idx
        elif last_mpn and idx - last_idx <= 3:
            mpn, inherited = last_mpn, True
            ctx = f"{last_ctx}  ⏎  {ctx}"
            last_idx = idx
        for r in unique(refs):
            data[r].append(SmtEntry(r, mpn or "", loc, ctx, inherited))
    return dict(data)


# =============================================================================
# Comparison
# =============================================================================
def compare_data(bom: BomData, smt: dict, scan_prefixes: set, opts: CompareOptions):
    ignore = {p.upper() for p in opts.ignore_prefixes}
    bom_refs = {r for r in bom.entries if ref_prefix(r) not in ignore}
    smt_refs = {r for r in smt if ref_prefix(r) not in ignore}

    rows = []
    stats = {t: 0 for t in ALL_TYPES}
    verified = unverified = 0

    def add(ref, typ, bom_mpn="", smt_mpn="", notes="", desc="", src=""):
        rows.append({COL_REF: ref, COL_TYPE: typ, COL_BOM_MPN: bom_mpn, COL_SMT_MPN: smt_mpn,
                     COL_NOTES: notes, COL_BOM_DESC: desc, COL_SMT_SRC: src})
        stats[typ] += 1

    for ref in sorted(bom_refs | smt_refs, key=natural_key):
        b = bom.entries.get(ref, []) if ref in bom_refs else []
        s = smt.get(ref, []) if ref in smt_refs else []
        b_mpn = "; ".join(unique(e.mpn for e in b))
        b_desc = "; ".join(unique(e.desc for e in b))
        s_mpn = "; ".join(unique(e.mpn for e in s))
        s_src = "; ".join(unique(e.location for e in s))

        if len(b) > 1:
            distinct = len({normalize_mpn(e.mpn) for e in b})
            note = f"Listed {len(b)} times in BOM (rows {', '.join(str(e.row) for e in b)})"
            note += " with DIFFERENT part numbers." if distinct > 1 else "."
            add(ref, DUP_BOM, b_mpn, s_mpn, note, b_desc, s_src)

        s_norms = {normalize_mpn(e.mpn) for e in s if e.mpn}
        if len(s_norms) > 1:
            add(ref, DUP_SMT, b_mpn, s_mpn,
                f"Designator appears on {len(s)} setup entries with different parts.", b_desc, s_src)

        if b and not s:
            note = "Present in BOM but not found on the SMT setup sheet."
            if ref_prefix(ref) not in scan_prefixes:
                note += f" Prefix '{ref_prefix(ref)}' is not in the SMT scan list."
            if any(e.inherited for e in b):
                note += " (BOM MPN inherited from previous row.)"
            add(ref, MISSING_SMT, b_mpn, "", note, b_desc, "")
        elif s and not b:
            if ref in bom.dnp_refs:
                note = f"Marked DNP / not fitted in BOM (row {bom.dnp_refs[ref]}) but loaded on the SMT setup."
            else:
                note = "Found on the SMT setup sheet but not in the BOM."
            add(ref, MISSING_BOM, "", s_mpn, note, "", s_src)
        elif b and s and opts.check_mpn:
            bom_norms = [n for n in (normalize_mpn(e.mpn) for e in b) if n]
            if not bom_norms:
                continue
            desc_norms = [n for n in (normalize_mpn(e.desc) for e in b) if len(n) >= 6]
            ok = False
            for e in s:
                ctx_n, sm_n = normalize_mpn(e.context), normalize_mpn(e.mpn)
                if any((len(n) >= 4 and n in ctx_n) or (sm_n and n == sm_n) for n in bom_norms):
                    ok = True
                elif any(d in ctx_n for d in desc_norms):
                    ok = True
                if ok:
                    break
            if ok:
                verified += 1
            elif s_mpn:
                note = f"BOM MPN not found on the setup line(s) at {s_src}; setup sheet shows '{s_mpn}'."
                if any(e.inherited for e in s):
                    note += " (SMT part number taken from the preceding line.)"
                add(ref, MPN_MISMATCH, b_mpn, s_mpn, note, b_desc, s_src)
            else:
                unverified += 1

    stats.update({
        "bom_refs": len(bom_refs),
        "smt_refs": len(smt_refs),
        "common": len(bom_refs & smt_refs),
        "mpn_verified": verified,
        "mpn_unverified": unverified,
        "dnp_excluded": len(bom.dnp_refs),
        "total": len(rows),
    })
    return rows, stats


def run_comparison(opts: CompareOptions, report: Callable[[int, str], None]) -> CompareResult:
    report(5, "Reading BOM…")
    bom = load_bom(opts)

    report(35, f"BOM loaded ({len(bom.entries)} designators). Reading SMT setup sheet…")
    lines = read_smt_lines(opts.smt_path)

    report(60, "Scanning setup sheet for designators…")
    prefixes = set(opts.prefixes)
    if opts.learn_prefixes:
        prefixes |= {ref_prefix(r) for r in list(bom.entries) + list(bom.dnp_refs)}
    scanner = SmtScanner(prefixes)
    bom_mpn_norms = {n for es in bom.entries.values() for e in es
                     for n in [normalize_mpn(e.mpn)] if len(n) >= 4}
    if opts.smt_column_mode:
        report(62, f"Reading setup columns {opts.smt_ref_col} (designators) / {opts.smt_mpn_col} (MPN)…")
        smt = parse_smt_columns(lines, scanner, bom_mpn_norms, opts.smt_ref_col, opts.smt_mpn_col)
    else:
        smt = parse_smt(lines, scanner, bom_mpn_norms)
    if not smt:
        extra = ""
        if opts.smt_column_mode:
            extra = (f"\n\nColumn mode is on (designators in column {opts.smt_ref_col}, "
                     f"MPN in column {opts.smt_mpn_col}). Turn it off for free-form sheets, "
                     "or change the column numbers.")
        raise ParseError(
            "No reference designators were found in the SMT setup sheet.\n\n"
            "Check that the correct file was selected and that the designator prefixes "
            "(e.g. R, C, U) match the ones used on the sheet." + extra
        )

    report(85, "Comparing…")
    rows, stats = compare_data(bom, smt, scanner.prefixes, opts)

    warnings = []
    if bom.invalid_tokens:
        sample = "\n".join(f"  Row {r}: '{t}'" for r, t in bom.invalid_tokens[:40])
        more = f"\n  … and {len(bom.invalid_tokens) - 40} more" if len(bom.invalid_tokens) > 40 else ""
        warnings.append(f"{len(bom.invalid_tokens)} token(s) in the BOM designator column were not "
                        f"recognised as designators and were ignored:\n{sample}{more}")
    report(100, "Done")
    return CompareResult(rows, stats, warnings, bom.header_row)


# =============================================================================
# Excel export
# =============================================================================
def export_report(path: str, rows: list, stats: dict, meta: dict):
    try:
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise FileReadError("Excel export requires 'openpyxl'.\n\nInstall with:  pip install openpyxl")

    df = pd.DataFrame(rows, columns=EXPORT_COLUMNS)
    summary = [
        ("Report generated", _dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("BOM file", meta.get("bom", "")),
        ("SMT setup sheet", meta.get("smt", "")),
        ("Rows exported", meta.get("scope", "")),
        ("", ""),
        ("Designators in BOM", stats.get("bom_refs", 0)),
        ("Designators on SMT sheet", stats.get("smt_refs", 0)),
        ("Designators in both", stats.get("common", 0)),
        ("MPN verified", stats.get("mpn_verified", 0)),
        ("MPN not verifiable (no part number on setup line)", stats.get("mpn_unverified", 0)),
        ("DNP designators excluded from BOM", stats.get("dnp_excluded", 0)),
        ("", ""),
    ] + [(t, stats.get(t, 0)) for t in ALL_TYPES] + [("Total discrepancies", stats.get("total", 0))]
    summary_df = pd.DataFrame(summary, columns=["Item", "Value"])

    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")

    def style_sheet(ws, max_width=70, type_col=None):
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
        ws.freeze_panes = "A2"
        if ws.max_row > 1:
            ws.auto_filter.ref = ws.dimensions
        for idx, column in enumerate(ws.columns, 1):
            longest = max((len(str(c.value)) for c in column if c.value is not None), default=8)
            ws.column_dimensions[get_column_letter(idx)].width = min(max(12, longest + 2), max_width)
        for row in ws.iter_rows(min_row=2):
            fill = None
            if type_col is not None:
                color = TYPE_COLORS.get(row[type_col].value)
                if color:
                    fill = PatternFill("solid", fgColor=color.lstrip("#"))
            for c in row:
                c.border = border
                c.alignment = Alignment(vertical="top", wrap_text=True)
                if fill:
                    c.fill = fill

    try:
        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            summary_df.to_excel(writer, sheet_name="Summary", index=False)
            df.to_excel(writer, sheet_name="Discrepancies", index=False)
            style_sheet(writer.sheets["Summary"], max_width=90)
            style_sheet(writer.sheets["Discrepancies"], type_col=EXPORT_COLUMNS.index(COL_TYPE))
    except PermissionError:
        raise FileReadError(f"Cannot write to:\n{path}\n\nThe file may be open in Excel. Close it and try again.")
    except OSError as e:
        raise FileReadError(f"Failed to save the report.\n\nDetails: {e}")


# =============================================================================
# Worker thread
# =============================================================================
class CompareWorker(QObject):
    progress = Signal(int, str)
    succeeded = Signal(object)
    failed = Signal(str, str, str)  # title, message, details

    def __init__(self, opts: CompareOptions):
        super().__init__()
        self.opts = opts

    @Slot()
    def run(self):
        try:
            result = run_comparison(self.opts, self.progress.emit)
            self.succeeded.emit(result)
        except AppError as e:
            self.failed.emit(e.title, str(e), "")
        except MemoryError:
            self.failed.emit("Out of Memory", "The files are too large to process in available memory.", "")
        except Exception as e:  # noqa: BLE001
            self.failed.emit("Unexpected Error", f"An unexpected error occurred:\n{e}", traceback.format_exc())


# =============================================================================
# Table model / filter
# =============================================================================
class DiscrepancyModel(QAbstractTableModel):
    def __init__(self, parent=None):
        super().__init__(parent)
        self._rows: list = []
        self._dark = False

    def set_dark(self, dark: bool):
        self._dark = bool(dark)
        if self._rows:
            top = self.index(0, 0)
            bot = self.index(self.rowCount() - 1, self.columnCount() - 1)
            self.dataChanged.emit(top, bot, [Qt.ItemDataRole.BackgroundRole, Qt.ItemDataRole.ForegroundRole])

    def set_rows(self, rows: list):
        self.beginResetModel()
        self._rows = list(rows)
        self.endResetModel()

    def row_at(self, r: int) -> dict:
        return self._rows[r]

    def rowCount(self, parent=QModelIndex()):
        return 0 if parent.isValid() else len(self._rows)

    def columnCount(self, parent=QModelIndex()):
        return 0 if parent.isValid() else len(TABLE_COLUMNS)

    def data(self, index, role=Qt.ItemDataRole.DisplayRole):
        if not index.isValid():
            return None
        row = self._rows[index.row()]
        col = TABLE_COLUMNS[index.column()]
        value = row.get(col, "")
        if role == Qt.ItemDataRole.DisplayRole:
            return value
        if role == Qt.ItemDataRole.ToolTipRole:
            tip = f"{col}: {value}"
            if row.get(COL_BOM_DESC):
                tip += f"\nBOM description: {row[COL_BOM_DESC]}"
            if row.get(COL_SMT_SRC):
                tip += f"\nSMT source: {row[COL_SMT_SRC]}"
            return tip
        if role == Qt.ItemDataRole.BackgroundRole:
            palette = DARK_TYPE_COLORS if self._dark else TYPE_COLORS
            color = palette.get(row.get(COL_TYPE))
            if not color:
                return None
            qc = QColor(color)
            if index.row() % 2:
                qc = qc.darker(108)
            return qc
        if role == Qt.ItemDataRole.ForegroundRole:
            return QColor("#E8EEF4") if self._dark else QColor("#1F2933")
        if role == Qt.ItemDataRole.FontRole and col == COL_TYPE:
            f = QFont()
            f.setBold(True)
            return f
        if role == SORT_ROLE:
            return natural_sort_string(value) if col == COL_REF else str(value).lower()
        return None

    def headerData(self, section, orientation, role=Qt.ItemDataRole.DisplayRole):
        if role == Qt.ItemDataRole.DisplayRole and orientation == Qt.Orientation.Horizontal:
            return TABLE_COLUMNS[section]
        return None


class DiscrepancyFilter(QSortFilterProxyModel):
    ALL = "All types"

    def __init__(self, parent=None):
        super().__init__(parent)
        self._terms: list = []
        self._type = self.ALL
        self.setSortRole(SORT_ROLE)

    def set_text(self, text: str):
        self._terms = [t for t in re.split(r"[,\s]+", text.strip().lower()) if t]
        self.invalidateFilter()

    def set_type(self, typ: str):
        self._type = typ
        self.invalidateFilter()

    def filterAcceptsRow(self, source_row, source_parent):
        row = self.sourceModel().row_at(source_row)
        if self._type != self.ALL and row.get(COL_TYPE) != self._type:
            return False
        if not self._terms:
            return True
        haystack = " ".join(str(row.get(c, "")) for c in EXPORT_COLUMNS).lower()
        # Any term matching shows the row (allows "R1 C5 U3" multi-lookup)
        return any(t in haystack for t in self._terms)


class CopyableTableView(QTableView):
    def keyPressEvent(self, event):
        if event.matches(QKeySequence.StandardKey.Copy):
            rows = sorted({i.row() for i in self.selectionModel().selectedIndexes()})
            model = self.model()
            lines = ["\t".join(TABLE_COLUMNS)]
            for r in rows:
                lines.append("\t".join(str(model.index(r, c).data() or "") for c in range(model.columnCount())))
            QApplication.clipboard().setText("\n".join(lines))
            return
        super().keyPressEvent(event)


class FileLineEdit(QLineEdit):
    """Line edit that accepts drag-and-drop of a single file with an allowed extension."""

    fileDropped = Signal(str)

    def __init__(self, extensions, placeholder, parent=None):
        super().__init__(parent)
        self._ext = tuple(e.lower() for e in extensions)
        self.setPlaceholderText(placeholder)
        self.setAcceptDrops(True)

    def _path_from(self, event) -> Optional[str]:
        md = event.mimeData()
        if md.hasUrls() and md.urls():
            p = md.urls()[0].toLocalFile()
            if p and p.lower().endswith(self._ext):
                return p
        return None

    def dragEnterEvent(self, event):
        if self._path_from(event):
            event.acceptProposedAction()
        else:
            event.ignore()

    def dragMoveEvent(self, event):
        self.dragEnterEvent(event)

    def dropEvent(self, event):
        p = self._path_from(event)
        if p:
            self.setText(p)
            self.fileDropped.emit(p)
            event.acceptProposedAction()


# =============================================================================
# Main window
# =============================================================================
class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle(f"{APP_NAME}  v{APP_VERSION}")
        self.resize(1400, 860)
        self.settings = QSettings(APP_ORG, "BomSmtComparator")
        self._thread: Optional[QThread] = None
        self._worker: Optional[CompareWorker] = None
        self._result: Optional[CompareResult] = None
        self._loading_headers = False
        self._dark = False
        self._build_ui()
        self._restore_settings()
        self._apply_theme(self._dark)
        for key in ("Ctrl+Return", "Ctrl+Enter"):
            shortcut = QShortcut(QKeySequence(key), self)
            shortcut.activated.connect(self._start_compare)
        export_shortcut = QShortcut(QKeySequence("Ctrl+E"), self)
        export_shortcut.activated.connect(self._export)

    # ------------------------------------------------------------------ UI build
    def _build_ui(self):
        splitter = QSplitter(Qt.Orientation.Horizontal)
        splitter.addWidget(self._build_left_panel())
        splitter.addWidget(self._build_results_panel())
        splitter.setStretchFactor(0, 0)
        splitter.setStretchFactor(1, 1)
        splitter.setSizes([430, 970])
        self.setCentralWidget(splitter)
        self.statusBar().showMessage("Select a setup sheet and a BOM to begin.")

    def _build_left_panel(self) -> QWidget:
        panel = QWidget()
        layout = QVBoxLayout(panel)
        layout.setContentsMargins(14, 14, 8, 14)
        layout.setSpacing(12)

        title_row = QHBoxLayout()
        title = QLabel("BOM ⇄ SMT Setup Comparator")
        title.setObjectName("title")
        title_row.addWidget(title, 1)
        self.theme_btn = QPushButton("Dark mode")
        self.theme_btn.setObjectName("theme")
        self.theme_btn.setCheckable(True)
        self.theme_btn.setToolTip("Toggle a darker shop-floor theme. Saved for next launch.")
        self.theme_btn.toggled.connect(self._on_theme_toggled)
        title_row.addWidget(self.theme_btn)
        layout.addLayout(title_row)
        subtitle = QLabel("Verify that what is loaded on the line matches the engineering BOM.")
        subtitle.setObjectName("subtitle")
        subtitle.setWordWrap(True)
        layout.addWidget(subtitle)

        # --- Files
        files_box = QGroupBox("1 · Input Files")
        fl = QVBoxLayout(files_box)
        fl.addWidget(QLabel("SMT Setup Sheet (.txt, .log, .pdf)"))
        self.smt_edit = FileLineEdit(SMT_EXTENSIONS, "Browse or drop a setup sheet…")
        self.smt_edit.fileDropped.connect(self._on_smt_selected)
        smt_btn = QPushButton("Browse…")
        smt_btn.clicked.connect(self._browse_smt)
        row = QHBoxLayout()
        row.addWidget(self.smt_edit, 1)
        row.addWidget(smt_btn)
        fl.addLayout(row)

        fl.addWidget(QLabel("Bill of Materials (.xlsx, .xls)"))
        self.bom_edit = FileLineEdit(BOM_EXTENSIONS, "Browse or drop a BOM workbook…")
        self.bom_edit.fileDropped.connect(self._on_bom_selected)
        self.bom_edit.editingFinished.connect(lambda: self._on_bom_selected(self.bom_edit.text().strip(), quiet=True))
        bom_btn = QPushButton("Browse…")
        bom_btn.clicked.connect(self._browse_bom)
        row = QHBoxLayout()
        row.addWidget(self.bom_edit, 1)
        row.addWidget(bom_btn)
        fl.addLayout(row)
        layout.addWidget(files_box)

        # --- BOM configuration
        cfg_box = QGroupBox("2 · BOM Columns")
        form = QFormLayout(cfg_box)
        form.setLabelAlignment(Qt.AlignmentFlag.AlignRight)
        self.sheet_combo = QComboBox()
        self.sheet_combo.currentIndexChanged.connect(lambda _: self._load_headers(quiet=False))
        self.header_spin = QSpinBox()
        self.header_spin.setRange(0, 10000)
        self.header_spin.setSpecialValueText("Auto-detect")
        self.header_spin.setToolTip("Excel row number that contains the column headers (0 = auto-detect).")
        self.header_spin.valueChanged.connect(lambda _: self._load_headers(quiet=False))

        self.ref_combo = self._make_column_combo("Header name, column number (1-based) or letter")
        self.mpn_combo = self._make_column_combo("Header name, column number (1-based) or letter")
        self.desc_combo = self._make_column_combo("Optional")
        self.desc_combo.addItem(NONE_OPTION)

        form.addRow("Sheet:", self.sheet_combo)
        form.addRow("Header row:", self.header_spin)
        form.addRow("Ref. Designators*:", self.ref_combo)
        form.addRow("MPN*:", self.mpn_combo)
        form.addRow("Description:", self.desc_combo)
        layout.addWidget(cfg_box)

        # --- Options
        opt_box = QGroupBox("3 · Comparison Options")
        ol = QFormLayout(opt_box)
        self.prefix_edit = QLineEdit(DEFAULT_PREFIXES)
        self.prefix_edit.setToolTip("Designator prefixes to look for in the setup sheet (comma separated).")
        self.ignore_edit = QLineEdit()
        self.ignore_edit.setPlaceholderText("e.g. TP, MH, FID, J")
        self.ignore_edit.setToolTip("Designators with these prefixes are excluded from the comparison "
                                    "(useful for through-hole, test points, fiducials…).")
        ol.addRow("SMT prefixes:", self.prefix_edit)
        ol.addRow("Ignore prefixes:", self.ignore_edit)
        self.learn_chk = QCheckBox("Also scan for prefixes used in the BOM")
        self.learn_chk.setChecked(True)
        self.mpn_chk = QCheckBox("Check MPN association (mismatch detection)")
        self.mpn_chk.setChecked(True)
        self.dnp_chk = QCheckBox("Exclude BOM rows marked DNP / not fitted")
        self.dnp_chk.setChecked(True)
        ol.addRow(self.learn_chk)
        ol.addRow(self.mpn_chk)
        ol.addRow(self.dnp_chk)
        self.col_mode_chk = QCheckBox("Setup sheet is columnar (tab / 2+ spaces / |)")
        self.col_mode_chk.setToolTip(
            "Use this only for machine dumps with a fixed column for designators. "
            "Leave off for typical feeder sheets."
        )
        self.col_mode_chk.toggled.connect(self._sync_column_mode)
        self.smt_ref_spin = QSpinBox()
        self.smt_ref_spin.setRange(1, 40)
        self.smt_ref_spin.setValue(6)
        self.smt_ref_spin.setToolTip("1-based column that holds reference designators.")
        self.smt_mpn_spin = QSpinBox()
        self.smt_mpn_spin.setRange(0, 40)
        self.smt_mpn_spin.setValue(3)
        self.smt_mpn_spin.setSpecialValueText("Guess")
        self.smt_mpn_spin.setToolTip("1-based MPN column. 0 = guess from the rest of the line.")
        ol.addRow(self.col_mode_chk)
        ol.addRow("Setup RefDes column:", self.smt_ref_spin)
        ol.addRow("Setup MPN column:", self.smt_mpn_spin)
        layout.addWidget(opt_box)

        # --- Action
        self.compare_btn = QPushButton("Compare Files")
        self.compare_btn.setObjectName("primary")
        self.compare_btn.setMinimumHeight(46)
        self.compare_btn.setToolTip("Ctrl+Enter")
        self.compare_btn.clicked.connect(self._start_compare)
        layout.addWidget(self.compare_btn)
        hint = QLabel("PDF must be selectable text, not a scan. Header Auto-detect is usually enough. Ctrl+Enter compares · Ctrl+E exports.")
        hint.setObjectName("subtitle")
        hint.setWordWrap(True)
        layout.addWidget(hint)

        self.progress = QProgressBar()
        self.progress.setRange(0, 100)
        self.progress.setTextVisible(True)
        self.progress.setVisible(False)
        layout.addWidget(self.progress)

        self.progress_label = QLabel("")
        self.progress_label.setObjectName("subtitle")
        self.progress_label.setWordWrap(True)
        layout.addWidget(self.progress_label)
        layout.addStretch(1)

        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setFrameShape(QFrame.Shape.NoFrame)
        scroll.setWidget(panel)
        scroll.setMinimumWidth(380)
        return scroll

    def _make_column_combo(self, placeholder: str) -> QComboBox:
        cb = QComboBox()
        cb.setEditable(True)
        cb.setInsertPolicy(QComboBox.InsertPolicy.NoInsert)
        cb.lineEdit().setPlaceholderText(placeholder)
        cb.setMinimumContentsLength(18)
        return cb

    def _build_results_panel(self) -> QWidget:
        panel = QWidget()
        layout = QVBoxLayout(panel)
        layout.setContentsMargins(8, 14, 14, 14)
        layout.setSpacing(10)

        header = QLabel("Discrepancies")
        header.setObjectName("section")
        layout.addWidget(header)

        self.summary_label = QLabel("No comparison run yet.")
        self.summary_label.setObjectName("summary")
        self.summary_label.setWordWrap(True)
        layout.addWidget(self.summary_label)

        self.chip_bar = QHBoxLayout()
        self.chip_bar.setSpacing(6)
        layout.addLayout(self.chip_bar)

        bar = QHBoxLayout()
        self.search_edit = QLineEdit()
        self.search_edit.setPlaceholderText("🔍  Search designators, MPNs, notes…  (space/comma = multiple terms)")
        self.search_edit.setClearButtonEnabled(True)
        self.type_combo = QComboBox()
        self.type_combo.addItems([DiscrepancyFilter.ALL] + ALL_TYPES)
        self.export_btn = QPushButton("Export to Excel")
        self.export_btn.setObjectName("secondary")
        self.export_btn.setToolTip("Ctrl+E")
        self.export_btn.setEnabled(False)
        self.export_btn.clicked.connect(self._export)
        bar.addWidget(self.search_edit, 1)
        bar.addWidget(self.type_combo)
        bar.addWidget(self.export_btn)
        layout.addLayout(bar)

        self.model = DiscrepancyModel(self)
        self.proxy = DiscrepancyFilter(self)
        self.proxy.setSourceModel(self.model)
        self.search_edit.textChanged.connect(self.proxy.set_text)
        self.search_edit.textChanged.connect(lambda _: self._update_count())
        self.type_combo.currentTextChanged.connect(self.proxy.set_type)
        self.type_combo.currentTextChanged.connect(lambda _: self._update_count())

        self.table = CopyableTableView()
        self.table.setModel(self.proxy)
        self.table.setSortingEnabled(True)
        self.table.sortByColumn(0, Qt.SortOrder.AscendingOrder)
        self.table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.table.setWordWrap(False)
        self.table.setAlternatingRowColors(True)
        self.table.verticalHeader().setVisible(False)
        self.table.verticalHeader().setDefaultSectionSize(26)
        hh = self.table.horizontalHeader()
        hh.setSectionResizeMode(QHeaderView.ResizeMode.Interactive)
        hh.setStretchLastSection(True)
        for i, w in enumerate((150, 150, 210, 210)):
            self.table.setColumnWidth(i, w)
        layout.addWidget(self.table, 1)

        legend = QHBoxLayout()
        for label, typ in (("Missing", MISSING_SMT), ("Mismatch", MPN_MISMATCH), ("Duplicate", DUP_BOM)):
            chip = QPushButton(label)
            chip.setCursor(Qt.CursorShape.PointingHandCursor)
            chip.setStyleSheet(
                f"QPushButton {{ background:{TYPE_COLORS[typ]}; color:#1F2933; border:none; "
                f"border-radius:4px; padding:3px 8px; }}"
            )
            chip.clicked.connect(lambda _, t=typ: self._filter_type(t))
            legend.addWidget(chip)
        legend.addStretch(1)
        self.count_label = QLabel("")
        self.count_label.setObjectName("subtitle")
        legend.addWidget(self.count_label)
        layout.addLayout(legend)
        return panel

    # --------------------------------------------------------------- settings
    def _restore_settings(self):
        s = self.settings
        geo = s.value("geometry")
        if geo is not None:
            self.restoreGeometry(geo)
        self.prefix_edit.setText(s.value("prefixes", DEFAULT_PREFIXES))
        self.ignore_edit.setText(s.value("ignore", ""))
        self.learn_chk.setChecked(s.value("learn", True, type=bool))
        self.mpn_chk.setChecked(s.value("check_mpn", True, type=bool))
        self.dnp_chk.setChecked(s.value("exclude_dnp", True, type=bool))
        self.col_mode_chk.setChecked(s.value("smt_column_mode", False, type=bool))
        self.smt_ref_spin.setValue(int(s.value("smt_ref_col", 6)))
        self.smt_mpn_spin.setValue(int(s.value("smt_mpn_col", 3)))
        self._dark = s.value("dark", False, type=bool)
        self.theme_btn.blockSignals(True)
        self.theme_btn.setChecked(self._dark)
        self.theme_btn.blockSignals(False)
        self._sync_column_mode(self.col_mode_chk.isChecked())

    def _save_settings(self):
        s = self.settings
        s.setValue("geometry", self.saveGeometry())
        s.setValue("prefixes", self.prefix_edit.text())
        s.setValue("ignore", self.ignore_edit.text())
        s.setValue("learn", self.learn_chk.isChecked())
        s.setValue("check_mpn", self.mpn_chk.isChecked())
        s.setValue("exclude_dnp", self.dnp_chk.isChecked())
        s.setValue("smt_column_mode", self.col_mode_chk.isChecked())
        s.setValue("smt_ref_col", self.smt_ref_spin.value())
        s.setValue("smt_mpn_col", self.smt_mpn_spin.value())
        s.setValue("dark", self._dark)

    def _last_dir(self) -> str:
        return self.settings.value("last_dir", os.path.expanduser("~"))

    def _remember_dir(self, path: str):
        self.settings.setValue("last_dir", os.path.dirname(path))

    # ------------------------------------------------------------- file picks
    def _browse_smt(self):
        path, _ = QFileDialog.getOpenFileName(
            self, "Select SMT Setup Sheet", self._last_dir(),
            "Setup sheets (*.txt *.log *.pdf);;Text files (*.txt *.log);;PDF files (*.pdf);;All files (*)")
        if path:
            self._remember_dir(path)
            self.smt_edit.setText(path)
            self._on_smt_selected(path)

    def _on_smt_selected(self, path: str):
        if not path or not os.path.isfile(path):
            return
        try:
            lines = read_smt_lines(path)
        except AppError as e:
            self.statusBar().showMessage(str(e).split("\n", 1)[0])
            return
        n = sum(1 for _, text in lines if text.strip())
        self.statusBar().showMessage(f"Setup sheet loaded: {n} text line(s).")

    def _browse_bom(self):
        path, _ = QFileDialog.getOpenFileName(
            self, "Select Bill of Materials", self._last_dir(),
            "Excel workbooks (*.xlsx *.xls *.xlsm);;All files (*)")
        if path:
            self._remember_dir(path)
            self.bom_edit.setText(path)
            self._on_bom_selected(path)

    def _on_bom_selected(self, path: str, quiet: bool = False):
        if not path or not os.path.isfile(path) or not path.lower().endswith(BOM_EXTENSIONS):
            if not quiet and path:
                self._warn("Invalid BOM", "Please select an existing .xlsx or .xls file.")
            return
        if getattr(self, "_loaded_bom_path", None) == path and quiet:
            return
        self._loaded_bom_path = path
        QApplication.setOverrideCursor(Qt.CursorShape.WaitCursor)
        try:
            sheets = list_bom_sheets(path)
        except AppError as e:
            QApplication.restoreOverrideCursor()
            self._error(e.title, str(e))
            return
        QApplication.restoreOverrideCursor()

        self._loading_headers = True
        self.sheet_combo.clear()
        self.sheet_combo.addItems(sheets)
        # Prefer a sheet that looks like a BOM
        for i, name in enumerate(sheets):
            if "bom" in name.lower():
                self.sheet_combo.setCurrentIndex(i)
                break
        self.header_spin.setValue(0)
        self._loading_headers = False
        self._load_headers(quiet=False)

    def _load_headers(self, quiet: bool = True):
        if self._loading_headers:
            return
        path = self.bom_edit.text().strip()
        if not path or not os.path.isfile(path) or self.sheet_combo.count() == 0:
            return
        QApplication.setOverrideCursor(Qt.CursorShape.WaitCursor)
        try:
            headers, idx = peek_bom_headers(path, self.sheet_combo.currentText(), self.header_spin.value())
        except AppError as e:
            QApplication.restoreOverrideCursor()
            if not quiet:
                self._error(e.title, str(e))
            return
        QApplication.restoreOverrideCursor()

        for combo, kind in ((self.ref_combo, "ref"), (self.mpn_combo, "mpn"), (self.desc_combo, "desc")):
            combo.blockSignals(True)
            combo.clear()
            if kind == "desc":
                combo.addItem(NONE_OPTION)
            combo.addItems(headers)
            guess = guess_column(headers, kind)
            combo.setCurrentText(guess if guess else (NONE_OPTION if kind == "desc" else ""))
            combo.blockSignals(False)

        mode = "auto-detected" if self.header_spin.value() == 0 else "set"
        self.statusBar().showMessage(f"BOM headers loaded from row {idx + 1} ({mode}). {len(headers)} columns found.")

    # --------------------------------------------------------------- compare
    def _collect_options(self) -> Optional[CompareOptions]:
        smt = self.smt_edit.text().strip()
        bom = self.bom_edit.text().strip()
        if not smt:
            self._warn("Missing File", "Please select an SMT setup sheet.")
            return None
        if not bom:
            self._warn("Missing File", "Please select a BOM Excel file.")
            return None
        if not os.path.isfile(smt):
            self._warn("File Not Found", f"The SMT setup sheet does not exist:\n{smt}")
            return None
        if not os.path.isfile(bom):
            self._warn("File Not Found", f"The BOM file does not exist:\n{bom}")
            return None
        if not smt.lower().endswith(SMT_EXTENSIONS):
            self._warn("Unsupported File", "The SMT setup sheet must be a .txt, .log or .pdf file.")
            return None
        if not bom.lower().endswith(BOM_EXTENSIONS):
            self._warn("Unsupported File", "The BOM must be an .xlsx or .xls file.")
            return None
        if not self.ref_combo.currentText().strip():
            self._warn("Missing Column", "Please specify the Reference Designator column.")
            return None
        if not self.mpn_combo.currentText().strip():
            self._warn("Missing Column", "Please specify the Manufacturer Part Number (MPN) column.")
            return None
        prefixes = parse_prefix_list(self.prefix_edit.text())
        if not prefixes and not self.learn_chk.isChecked():
            self._warn("Missing Prefixes", "Enter at least one designator prefix (e.g. R, C, U).")
            return None
        return CompareOptions(
            smt_path=smt,
            bom_path=bom,
            bom_sheet=self.sheet_combo.currentText() or 0,
            header_row=self.header_spin.value(),
            ref_col=self.ref_combo.currentText(),
            mpn_col=self.mpn_combo.currentText(),
            desc_col=self.desc_combo.currentText(),
            prefixes=prefixes,
            ignore_prefixes=set(parse_prefix_list(self.ignore_edit.text())),
            learn_prefixes=self.learn_chk.isChecked(),
            check_mpn=self.mpn_chk.isChecked(),
            exclude_dnp=self.dnp_chk.isChecked(),
            smt_column_mode=self.col_mode_chk.isChecked(),
            smt_ref_col=self.smt_ref_spin.value(),
            smt_mpn_col=self.smt_mpn_spin.value(),
        )

    def _start_compare(self):
        if self._thread is not None:
            return
        opts = self._collect_options()
        if opts is None:
            return
        self._save_settings()
        self._set_busy(True)

        self._thread = QThread(self)
        self._worker = CompareWorker(opts)
        self._worker.moveToThread(self._thread)
        self._thread.started.connect(self._worker.run)
        self._worker.progress.connect(self._on_progress)
        self._worker.succeeded.connect(self._on_success)
        self._worker.failed.connect(self._on_failed)
        self._worker.succeeded.connect(self._thread.quit)
        self._worker.failed.connect(self._thread.quit)
        self._thread.finished.connect(self._cleanup_thread)
        self._thread.start()

    def _set_busy(self, busy: bool):
        for w in (self.compare_btn, self.smt_edit, self.bom_edit, self.sheet_combo, self.header_spin,
                  self.ref_combo, self.mpn_combo, self.desc_combo, self.prefix_edit, self.ignore_edit,
                  self.col_mode_chk, self.smt_ref_spin, self.smt_mpn_spin, self.theme_btn):
            w.setEnabled(not busy)
        if not busy:
            self._sync_column_mode(self.col_mode_chk.isChecked())
        self.compare_btn.setText("Comparing…" if busy else "Compare Files")
        self.progress.setVisible(busy)
        if busy:
            self.progress.setValue(0)
            QApplication.setOverrideCursor(Qt.CursorShape.BusyCursor)
        else:
            QApplication.restoreOverrideCursor()
            self.progress_label.setText("")

    @Slot(int, str)
    def _on_progress(self, value: int, message: str):
        self.progress.setValue(value)
        self.progress_label.setText(message)
        self.statusBar().showMessage(message)

    @Slot(object)
    def _on_success(self, result: CompareResult):
        self._set_busy(False)
        self._result = result
        self.model.set_rows(result.rows)
        self.table.sortByColumn(0, Qt.SortOrder.AscendingOrder)
        self.export_btn.setEnabled(True)
        self._refresh_type_filter(result)
        self._update_summary(result)
        self._update_count()
        st = result.stats
        self.statusBar().showMessage(
            f"Comparison complete: {st['total']} discrepancies "
            f"(BOM header row {result.header_row})."
        )
        if result.warnings:
            box = QMessageBox(self)
            box.setIcon(QMessageBox.Icon.Information)
            box.setWindowTitle("Comparison Finished With Warnings")
            box.setText("The comparison finished, but some BOM data could not be interpreted.\n"
                        "Click 'Show Details…' to review.")
            box.setDetailedText("\n\n".join(result.warnings))
            box.exec()

    @Slot(str, str, str)
    def _on_failed(self, title: str, message: str, details: str):
        self._set_busy(False)
        self.statusBar().showMessage("Comparison failed.")
        self._error(title, message, details)

    @Slot()
    def _cleanup_thread(self):
        if self._worker is not None:
            self._worker.deleteLater()
        if self._thread is not None:
            self._thread.deleteLater()
        self._worker = None
        self._thread = None

    def _update_summary(self, result: CompareResult):
        st = result.stats
        if st["total"] == 0:
            self.summary_label.setText(
                f"BOM: {st['bom_refs']} designators  ·  SMT: {st['smt_refs']}  ·  "
                f"In both: {st['common']}  ·  MPN verified: {st['mpn_verified']}  ·  "
                f"DNP excluded: {st['dnp_excluded']}\n"
                "✔ The BOM and SMT setup sheet match. No discrepancies were found."
            )
        else:
            self.summary_label.setText(
                f"BOM: {st['bom_refs']} designators  ·  SMT: {st['smt_refs']}  ·  "
                f"In both: {st['common']}  ·  MPN verified: {st['mpn_verified']}  ·  "
                f"Unverifiable: {st['mpn_unverified']}  ·  DNP excluded: {st['dnp_excluded']}"
            )
        self._rebuild_chips(st)

    def _clear_layout(self, layout: QHBoxLayout):
        while layout.count():
            item = layout.takeAt(0)
            widget = item.widget()
            if widget is not None:
                widget.deleteLater()

    def _chip_button(self, text: str, color: Optional[str], handler) -> QPushButton:
        btn = QPushButton(text)
        btn.setCursor(Qt.CursorShape.PointingHandCursor)
        if color:
            btn.setStyleSheet(
                f"QPushButton {{ background:{color}; color:#1F2933; border:none; "
                f"border-radius:6px; padding:4px 10px; font-weight:600; }}"
            )
        btn.clicked.connect(handler)
        return btn

    def _rebuild_chips(self, st: dict):
        self._clear_layout(self.chip_bar)
        if not st.get("total"):
            banner = QLabel("✔ No discrepancies")
            banner.setObjectName("okBanner")
            self.chip_bar.addWidget(banner)
            self.chip_bar.addStretch(1)
            return
        self.chip_bar.addWidget(self._chip_button("All", None, lambda: self._filter_type(DiscrepancyFilter.ALL)))
        for typ in ALL_TYPES:
            count = st.get(typ) or 0
            if not count:
                continue
            self.chip_bar.addWidget(
                self._chip_button(f"{typ}: {count}", TYPE_COLORS[typ], lambda _, t=typ: self._filter_type(t))
            )
        self.chip_bar.addStretch(1)

    def _filter_type(self, typ: str):
        idx = self.type_combo.findText(typ)
        if idx >= 0:
            self.type_combo.setCurrentIndex(idx)

    def _refresh_type_filter(self, result: CompareResult):
        present = [t for t in ALL_TYPES if result.stats.get(t)]
        current = self.type_combo.currentText()
        self.type_combo.blockSignals(True)
        self.type_combo.clear()
        self.type_combo.addItem(DiscrepancyFilter.ALL)
        self.type_combo.addItems(present)
        if current in present or current == DiscrepancyFilter.ALL:
            self.type_combo.setCurrentText(current)
        else:
            self.type_combo.setCurrentText(DiscrepancyFilter.ALL)
        self.type_combo.blockSignals(False)
        self.proxy.set_type(self.type_combo.currentText())

    def _sync_column_mode(self, on: bool):
        self.smt_ref_spin.setEnabled(on)
        self.smt_mpn_spin.setEnabled(on)

    def _on_theme_toggled(self, dark: bool):
        self._dark = bool(dark)
        self._apply_theme(self._dark)
        self._save_settings()

    def _apply_theme(self, dark: bool):
        app = QApplication.instance()
        if app is None:
            return
        self.theme_btn.setText("Light mode" if dark else "Dark mode")
        if dark:
            apply_dark_palette(app)
            app.setStyleSheet(DARK_STYLE_SHEET)
        else:
            apply_light_palette(app)
            app.setStyleSheet(STYLE_SHEET)
        if hasattr(self, "model"):
            self.model.set_dark(dark)

    def _update_count(self):
        total = self.model.rowCount()
        shown = self.proxy.rowCount()
        self.count_label.setText(f"Showing {shown} of {total} rows" if total else "")

    # ---------------------------------------------------------------- export
    def _visible_rows(self) -> list:
        rows = []
        for r in range(self.proxy.rowCount()):
            src = self.proxy.mapToSource(self.proxy.index(r, 0))
            rows.append(self.model.row_at(src.row()))
        return rows

    def _export(self):
        if not self._result or self.model.rowCount() == 0:
            self._warn("Nothing to Export", "Run a comparison with discrepancies before exporting.")
            return
        rows = self._visible_rows()
        filtered = len(rows) != self.model.rowCount()
        if filtered:
            choice = QMessageBox.question(
                self, "Export Scope",
                f"A filter is active ({len(rows)} of {self.model.rowCount()} rows shown).\n\n"
                "Export only the visible rows?\n(Choose 'No' to export all rows.)",
                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No | QMessageBox.StandardButton.Cancel)
            if choice == QMessageBox.StandardButton.Cancel:
                return
            if choice == QMessageBox.StandardButton.No:
                rows, filtered = self.model._rows, False
        if not rows:
            self._warn("Nothing to Export", "No rows match the current filter.")
            return

        default_name = f"BOM_SMT_Discrepancies_{_dt.datetime.now():%Y%m%d_%H%M}.xlsx"
        path, _ = QFileDialog.getSaveFileName(self, "Export Discrepancy Report",
                                              os.path.join(self._last_dir(), default_name),
                                              "Excel workbook (*.xlsx)")
        if not path:
            return
        if not path.lower().endswith(".xlsx"):
            path += ".xlsx"
        meta = {"bom": self.bom_edit.text(), "smt": self.smt_edit.text(),
                "scope": f"{len(rows)} (filtered)" if filtered else f"{len(rows)} (all)"}
        QApplication.setOverrideCursor(Qt.CursorShape.WaitCursor)
        try:
            export_report(path, rows, self._result.stats, meta)
        except AppError as e:
            QApplication.restoreOverrideCursor()
            self._error(e.title, str(e))
            return
        except Exception as e:  # noqa: BLE001
            QApplication.restoreOverrideCursor()
            self._error("Export Failed", f"Could not export the report:\n{e}", traceback.format_exc())
            return
        QApplication.restoreOverrideCursor()
        self._remember_dir(path)
        self.statusBar().showMessage(f"Report saved: {path}")
        QMessageBox.information(self, "Export Complete", f"Report saved to:\n{path}")

    # --------------------------------------------------------------- dialogs
    def _warn(self, title: str, message: str):
        QMessageBox.warning(self, title, message)

    def _error(self, title: str, message: str, details: str = ""):
        box = QMessageBox(self)
        box.setIcon(QMessageBox.Icon.Critical)
        box.setWindowTitle(title)
        box.setText(message)
        if details:
            box.setDetailedText(details)
        box.exec()

    def closeEvent(self, event):
        self._save_settings()
        if self._thread is not None and self._thread.isRunning():
            self._thread.quit()
            self._thread.wait(3000)
        super().closeEvent(event)


# =============================================================================
# Styling & entry point
# =============================================================================
STYLE_SHEET = """
QWidget { font-size: 10pt; color: #1F2933; }
QMainWindow, QScrollArea, QScrollArea > QWidget > QWidget { background: #F4F6FA; }
QLabel#title { font-size: 16pt; font-weight: 700; color: #102A43; }
QLabel#section { font-size: 13pt; font-weight: 700; color: #102A43; }
QLabel#subtitle { color: #627D98; }
QLabel#summary { background: #FFFFFF; border: 1px solid #D9E2EC; border-radius: 8px; padding: 8px; }
QGroupBox { background: #FFFFFF; border: 1px solid #D9E2EC; border-radius: 8px;
            margin-top: 14px; padding: 12px 10px 10px 10px; font-weight: 600; }
QGroupBox::title { subcontrol-origin: margin; left: 10px; padding: 0 4px; color: #334E68; }
QLineEdit, QComboBox, QSpinBox { background: #FFFFFF; border: 1px solid #BCCCDC; border-radius: 5px;
            padding: 5px 6px; min-height: 20px; font-weight: 400; }
QLineEdit:focus, QComboBox:focus, QSpinBox:focus { border: 1px solid #2680C2; }
QLineEdit:disabled, QComboBox:disabled, QSpinBox:disabled { background: #F0F4F8; color: #9FB3C8; }
QPushButton { background: #FFFFFF; border: 1px solid #BCCCDC; border-radius: 5px; padding: 6px 14px; }
QPushButton:hover { background: #F0F4F8; }
QPushButton#primary { background: #2680C2; color: white; border: none; font-size: 12pt; font-weight: 700; border-radius: 8px; }
QPushButton#primary:hover { background: #186FAF; }
QPushButton#primary:disabled { background: #9FB3C8; }
QPushButton#secondary { background: #1E7D32; color: white; border: none; font-weight: 600; }
QPushButton#secondary:hover { background: #17692A; }
QPushButton#secondary:disabled { background: #B8C9B9; }
QProgressBar { border: 1px solid #BCCCDC; border-radius: 6px; background: #FFFFFF; text-align: center; height: 18px; }
QProgressBar::chunk { background: #2680C2; border-radius: 5px; }
QTableView { background: #FFFFFF; border: 1px solid #D9E2EC; border-radius: 6px; gridline-color: #E4E7EB;
             selection-background-color: #2680C2; selection-color: white; }
QHeaderView::section { background: #102A43; color: white; padding: 6px; border: none;
                       border-right: 1px solid #243B53; font-weight: 600; }
QStatusBar { background: #E4E7EB; }
QCheckBox { font-weight: 400; }
QPushButton#theme { padding: 4px 10px; font-weight: 600; }
QLabel#okBanner { color: #1E7D32; font-weight: 700; }
"""


DARK_STYLE_SHEET = """
QWidget { font-size: 10pt; color: #E8EEF4; }
QMainWindow, QScrollArea, QScrollArea > QWidget > QWidget { background: #1B2430; }
QLabel#title { font-size: 16pt; font-weight: 700; color: #F0F7FC; }
QLabel#section { font-size: 13pt; font-weight: 700; color: #F0F7FC; }
QLabel#subtitle { color: #9FB3C8; }
QLabel#summary { background: #243040; border: 1px solid #3D4F63; border-radius: 8px; padding: 8px; color: #E8EEF4; }
QLabel#okBanner { color: #5FE08A; font-weight: 700; }
QGroupBox { background: #243040; border: 1px solid #3D4F63; border-radius: 8px;
            margin-top: 14px; padding: 12px 10px 10px 10px; font-weight: 600; color: #E8EEF4; }
QGroupBox::title { subcontrol-origin: margin; left: 10px; padding: 0 4px; color: #B6C9DC; }
QLineEdit, QComboBox, QSpinBox { background: #151C26; border: 1px solid #506078; border-radius: 5px;
            padding: 5px 6px; min-height: 20px; font-weight: 400; color: #E8EEF4; }
QLineEdit:focus, QComboBox:focus, QSpinBox:focus { border: 1px solid #2680C2; }
QLineEdit:disabled, QComboBox:disabled, QSpinBox:disabled { background: #1B2430; color: #627D98; }
QPushButton { background: #2F4258; color: white; border: 1px solid #506078; border-radius: 5px; padding: 6px 14px; }
QPushButton:hover { background: #3A526C; }
QPushButton#primary { background: #2680C2; color: white; border: none; font-size: 12pt; font-weight: 700; border-radius: 8px; }
QPushButton#primary:hover { background: #186FAF; }
QPushButton#primary:disabled { background: #3D4F63; }
QPushButton#secondary { background: #1E7D32; color: white; border: none; font-weight: 600; }
QPushButton#secondary:hover { background: #17692A; }
QPushButton#secondary:disabled { background: #355244; }
QPushButton#theme { padding: 4px 10px; font-weight: 600; }
QProgressBar { border: 1px solid #506078; border-radius: 6px; background: #151C26; text-align: center; height: 18px; color: #E8EEF4; }
QProgressBar::chunk { background: #2680C2; border-radius: 5px; }
QTableView { background: #151C26; alternate-background-color: #1B2430; border: 1px solid #3D4F63; border-radius: 6px;
             gridline-color: #3D4F63; color: #E8EEF4; selection-background-color: #2680C2; selection-color: white; }
QHeaderView::section { background: #102A43; color: white; padding: 6px; border: none;
                       border-right: 1px solid #243B53; font-weight: 600; }
QStatusBar { background: #141C26; color: #9FB3C8; }
QCheckBox { font-weight: 400; color: #E8EEF4; }
QScrollArea { background: #1B2430; }
"""


def apply_light_palette(app: QApplication):
    pal = QPalette()
    pal.setColor(QPalette.ColorRole.Window, QColor("#F4F6FA"))
    pal.setColor(QPalette.ColorRole.WindowText, QColor("#1F2933"))
    pal.setColor(QPalette.ColorRole.Base, QColor("#FFFFFF"))
    pal.setColor(QPalette.ColorRole.AlternateBase, QColor("#F8FAFC"))
    pal.setColor(QPalette.ColorRole.Text, QColor("#1F2933"))
    pal.setColor(QPalette.ColorRole.Button, QColor("#FFFFFF"))
    pal.setColor(QPalette.ColorRole.ButtonText, QColor("#1F2933"))
    pal.setColor(QPalette.ColorRole.Highlight, QColor("#2680C2"))
    pal.setColor(QPalette.ColorRole.HighlightedText, QColor("#FFFFFF"))
    pal.setColor(QPalette.ColorRole.ToolTipBase, QColor("#FFFFFF"))
    pal.setColor(QPalette.ColorRole.ToolTipText, QColor("#1F2933"))
    app.setPalette(pal)


def apply_dark_palette(app: QApplication):
    pal = QPalette()
    pal.setColor(QPalette.ColorRole.Window, QColor("#1B2430"))
    pal.setColor(QPalette.ColorRole.WindowText, QColor("#E8EEF4"))
    pal.setColor(QPalette.ColorRole.Base, QColor("#151C26"))
    pal.setColor(QPalette.ColorRole.AlternateBase, QColor("#1B2430"))
    pal.setColor(QPalette.ColorRole.Text, QColor("#E8EEF4"))
    pal.setColor(QPalette.ColorRole.Button, QColor("#2F4258"))
    pal.setColor(QPalette.ColorRole.ButtonText, QColor("#FFFFFF"))
    pal.setColor(QPalette.ColorRole.Highlight, QColor("#2680C2"))
    pal.setColor(QPalette.ColorRole.HighlightedText, QColor("#FFFFFF"))
    pal.setColor(QPalette.ColorRole.ToolTipBase, QColor("#243040"))
    pal.setColor(QPalette.ColorRole.ToolTipText, QColor("#E8EEF4"))
    pal.setColor(QPalette.ColorRole.PlaceholderText, QColor("#9FB3C8"))
    app.setPalette(pal)


def install_exception_hook():
    """Show any uncaught exception in a dialog instead of silently crashing."""
    def hook(exc_type, exc, tb):
        details = "".join(traceback.format_exception(exc_type, exc, tb))
        sys.stderr.write(details)
        app = QApplication.instance()
        if app is not None:
            box = QMessageBox()
            box.setIcon(QMessageBox.Icon.Critical)
            box.setWindowTitle("Unexpected Error")
            box.setText(f"An unexpected error occurred:\n{exc}")
            box.setDetailedText(details)
            box.exec()
    sys.excepthook = hook


def main():
    app = QApplication(sys.argv)
    app.setApplicationName(APP_NAME)
    app.setOrganizationName(APP_ORG)
    app.setStyle("Fusion")
    apply_light_palette(app)
    app.setStyleSheet(STYLE_SHEET)
    install_exception_hook()
    window = MainWindow()
    window.show()
    sys.exit(app.exec())


if __name__ == "__main__":
    main()
